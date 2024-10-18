import os
import shutil
import re
from PyQt5.QtWidgets import QMessageBox
from Back_End import config
from openpyxl import load_workbook
from datetime import datetime
from docx import Document

class CoverLetterController:
    def __init__(self, main_window):
        self.main_window = main_window

    def create_cover_sheet(self):
        # Switch to message area
        self.main_window.display_area.setCurrentIndex(0)
        self.main_window.message_area.append("\nCreate Cover Sheet")

        # Check if we have the necessary data
        if not all([
            hasattr(self.main_window, 'project'), hasattr(self.main_window, 'client')
        ]):
            self.main_window.message_area.append("No project or client data available. Please retrieve projects first.")
            return

        # Access the models
        project = self.main_window.project
        client = self.main_window.client

        # Normalize the project code
        project_code = project.code
        try:
            normalized_code, job_number, year_suffix = self.normalize_project_code(project_code)
        except ValueError as e:
            self.main_window.message_area.append(f"Error normalizing project code: {e}")
            return

        # Locate the project directory
        job_dir = self.find_project_directory(normalized_code, job_number, year_suffix)
        if not job_dir:
            self.main_window.message_area.append(f"Job directory not found for project code: {normalized_code}")
            return

        # Locate the Project_Reports directory
        reports_dir = os.path.join(job_dir, "Project_Reports")
        if not os.path.exists(reports_dir):
            self.main_window.message_area.append(f"Project_Reports directory does not exist in {job_dir}")
            return

        # Find the Excel report
        excel_files = [f for f in os.listdir(reports_dir) if f.endswith('.xlsx')]
        if not excel_files:
            self.main_window.message_area.append(f"No Excel report found in {reports_dir}")
            return
        excel_file_path = os.path.join(reports_dir, excel_files[0])  # Assuming the first one is the report

        # Extract data from the Excel workbook
        extracted_data = self.extract_data_from_excel(excel_file_path)

        if not extracted_data:
            self.main_window.message_area.append("Failed to extract data from the Excel report.")
            return

        # Copy and modify the Word template
        template_path = config.COVER_LETTER_TEMPLATE_PATH
        cover_letters_dir = config.COVER_LETTER_OUTPUT_DIR

        if not os.path.exists(template_path):
            self.main_window.message_area.append(f"Template file not found at {template_path}")
            return

        # Ensure the destination directory exists
        if not os.path.exists(cover_letters_dir):
            os.makedirs(cover_letters_dir)

        # Create the file name
        stripped_project_name = project.name.replace('/', ' ')
        cover_letter_filename = f"{normalized_code}_{stripped_project_name}_cover_letter.docx"
        cover_letter_path = os.path.join(cover_letters_dir, cover_letter_filename)

        # Copy the template
        shutil.copy(template_path, cover_letter_path)

        # Modify the copied template
        self.insert_data_into_template(cover_letter_path, extracted_data)

        self.main_window.message_area.append(f"Cover letter created and saved at {cover_letter_path}")

    def normalize_project_code(self, project_code):
        # Same as in ReportController
        parts = project_code.split('-')
        if len(parts) != 2:
            raise ValueError(f"Invalid project code format: {project_code}")
        part1, part2 = parts
        if len(part2) == 2 and len(part1) > 2:
            # Assume format is xxxx-yy
            job_number = part1
            year_suffix = part2
        elif len(part1) == 2 and len(part2) > 2:
            # Assume format is yy-xxxx, rearrange
            job_number = part2
            year_suffix = part1
        else:
            raise ValueError(f"Cannot determine format of project code: {project_code}")
        # Return normalized code in format xxxx-yy
        normalized_code = f"{job_number}-{year_suffix}"
        return normalized_code, job_number, year_suffix

    def find_project_directory(self, normalized_code, job_number, year_suffix):
        report_directory = config.REPORT_DIRECTORY
        year_directory = os.path.join(report_directory, f"Jobs 20{year_suffix}")

        if not os.path.isdir(year_directory):
            self.main_window.message_area.append(f"Year directory does not exist: {year_directory}")
            return None

        # Search for the job in the year directory
        for dir_name in os.listdir(year_directory):
            dir_path = os.path.join(year_directory, dir_name)
            if not os.path.isdir(dir_path):
                continue

            # Normalize codes for comparison
            dir_code = dir_name[:6].strip()
            # Remove non-digit characters and leading zeros
            dir_code_digits = re.sub(r'\D', '', dir_code).lstrip('0')
            normalized_code_digits = re.sub(r'\D', '', normalized_code).lstrip('0')

            if dir_code_digits == normalized_code_digits:
                return dir_path

        return None

    def extract_data_from_excel(self, excel_file_path):
        try:
            workbook = load_workbook(filename=excel_file_path)
            sheet = workbook.active

            # Extract data from specific cells
            extracted_data = {
                'Project': sheet['B2'].value,
                'Location': f"{sheet['C2'].value}, {sheet['D2'].value}, {sheet['E2'].value}, {sheet['F2'].value}, {sheet['G2'].value}",
                'Client': sheet['I2'].value,
                'Date': datetime.now().strftime('%B %d, %Y'),
                'Job No': sheet['A2'].value,
                'Code': sheet['B5'].value
            }

            return extracted_data
        except Exception as e:
            self.main_window.message_area.append(f"Error extracting data from Excel: {e}")
            return None

    def insert_data_into_template(self, docx_path, data):
        try:
            document = Document(docx_path)

            # Replace placeholders in the document
            for paragraph in document.paragraphs:
                self.replace_placeholder_in_paragraph(paragraph, data)

            # Also check headers, footers, tables, etc., if necessary
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            self.replace_placeholder_in_paragraph(paragraph, data)

            # Save the modified document
            document.save(docx_path)
        except Exception as e:
            self.main_window.message_area.append(f"Error modifying Word template: {e}")

    def replace_placeholder_in_paragraph(self, paragraph, data):
        for key, value in data.items():
            if f'{{{{{key}}}}}' in paragraph.text:
                inline = paragraph.runs
                for i in range(len(inline)):
                    if f'{{{{{key}}}}}' in inline[i].text:
                        text = inline[i].text.replace(f'{{{{{key}}}}}', str(value))
                        inline[i].text = text
